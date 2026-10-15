"""The native PowerPoint: 16:9; one slide per deck slide; real text (titles and bullets readable, bold leads); photos
as pictures cropped to fill; cards as rounded shapes; the chart as a native chart; speaker notes with the credits;
nothing outside the slide; every theme; a deck without photos; today's Markdown."""
import io

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

from odie.decks.model import parse
from odie.decks.photos import Photo
from odie.decks.pptx import build_pptx
from odie.decks.themes import THEMES, theme
from tests.test_decks_model import MEGHALAYA, STOCK


@pytest.fixture
def meg(tmp_path):
    p = tmp_path / "x.jpg"
    Image.new("RGB", (1600, 900), (10, 80, 50)).save(p, "JPEG")
    d = parse(MEGHALAYA)
    return d, {r: Photo(r, p, 1600, 900, c or "ex.com") for r, c in d.photo_refs()}


def texts(slide):
    return " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def pictures(slide):
    return [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]


def test_the_pptx_is_native_and_editable(meg):
    d, photos = meg
    prs = Presentation(io.BytesIO(build_pptx(d, theme("editorial"), photos)))
    assert len(prs.slides) == len(d.slides)
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.01
    cover, photo, steps, cards, stats, chart, gallery, split, closing = prs.slides
    assert "MEGHALAYA" in texts(cover) and pictures(cover)
    assert "Post-monsoon green" in texts(photo)
    bold = [r for sh in photo.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs if r.font.bold]
    assert any("Post-monsoon green" in r.text for r in bold)
    assert "Umiam → Shillong" in texts(steps)
    rounded = [sh for sh in cards.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE]
    assert len(rounded) >= 2 and "Ri Kynjai" in texts(cards) and len(pictures(cards)) == 1
    assert "₹60–80k" in texts(stats)
    assert any(sh.has_chart for sh in chart.shapes)
    assert len(pictures(gallery)) == 2 and len(pictures(split)) == 1
    assert "Thanks" in texts(closing)
    assert "Chalo Hoppo" in cover.notes_slide.notes_text_frame.text
    assert "Dawki is clearest" in photo.notes_slide.notes_text_frame.text


def test_photos_are_cropped_to_fill_their_box(meg):
    d, photos = meg  # 16:9 photos
    prs = Presentation(io.BytesIO(build_pptx(d, theme("editorial"), photos)))
    card_pic = pictures(prs.slides[3])[0]  # a card's photo box is wider than tall but not 16:9
    assert card_pic.crop_top + card_pic.crop_bottom > 0 or card_pic.crop_left + card_pic.crop_right > 0


@pytest.mark.parametrize("name", list(THEMES))
def test_nothing_falls_outside_the_slide(meg, name):
    d, photos = meg
    prs = Presentation(io.BytesIO(build_pptx(d, theme(name), photos)))
    for i, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            assert sh.left >= 0 and sh.top >= 0, (i, sh.name)
            assert sh.left + sh.width <= prs.slide_width + 2 and sh.top + sh.height <= prs.slide_height + 2, (i, sh.name)


def test_a_stock_deck_without_photos(tmp_path):
    prs = Presentation(io.BytesIO(build_pptx(parse(STOCK), theme("minimal"), {})))
    assert len(prs.slides) == 7 and not any(pictures(s) for s in prs.slides)
    assert any(sh.has_chart for sh in prs.slides[2].shapes) and "$46.7B" in texts(prs.slides[1])
    assert "The bull case" in texts(prs.slides[4]) and "Nov 19" in texts(prs.slides[5])


def test_old_markdown_still_builds():
    d = parse("# Home servers\nWhat a home server can do\n\n# Why now\n- Hardware is cheap\n> Mention 32 GB\n\n"
              "# Plans\n| Plan | Price |\n| --- | --- |\n| Mini | ₹60k |\n")
    prs = Presentation(io.BytesIO(build_pptx(d, theme(""), {})))
    assert len(prs.slides) == 3 and any(sh.has_table for sh in prs.slides[2].shapes)
    assert "Mention 32 GB" in prs.slides[1].notes_slide.notes_text_frame.text


def test_steps_and_stats_wear_their_photo(tmp_path):
    p = tmp_path / "x.jpg"
    Image.new("RGB", (1600, 900), (10, 80, 50)).save(p, "JPEG")
    d = parse("# A\n\n# Day 1\nimage: https://ex.com/d1.jpg\n1. a — b\n2. c — d\n3. e — f\n\n# Budget\nimage: https://ex.com/b.jpg\nstat: 1 — one\nstat: 2 — two\n")
    photos = {r: Photo(r, p, 1600, 900, "ex.com") for r, _ in d.photo_refs()}
    prs = Presentation(io.BytesIO(build_pptx(d, theme("minimal"), photos)))
    assert pictures(prs.slides[1]) and pictures(prs.slides[2])
    assert "Photo: ex.com" in prs.slides[1].notes_slide.notes_text_frame.text
