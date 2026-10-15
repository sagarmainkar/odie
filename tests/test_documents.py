"""The documents skill: Markdown the assistant writes becomes a real PowerPoint deck or Word document (built by Hermes
Agent's MIT-licensed builders, vendored), saved in the workspace and handed to the chat as a download."""
import pytest
from docx import Document
from pptx import Presentation

from odie import media
from odie.agent.tools.files import workspace
from odie.skills.documents import build
from odie.store.db import Store

DECK = """# Our podcast idea
A weekly show about home servers

# Why now
- Home servers are **cheap**
  - A home server is enough
- People want privacy
> Mention the 32 GB home server

# Costs
| Item | Price |
| --- | --- |
| home server | ₹60,000 |

# Thanks
"""

DOC = """# Trip plan
Day one in **Goa**, with *sun*.

## Packing
- Sunscreen
- Hat

1. Book a cab
2. Pack

| Item | Qty |
| --- | --- |
| Hat | 1 |
"""


@pytest.fixture
def store(isolated_settings, tmp_path):
    s = Store(tmp_path / "assistant.db")
    s.init_schema()
    return s


def test_slides_from_markdown():
    slides = build.slides_from_markdown(DECK)
    assert slides[0] == {"layout": "title", "title": "Our podcast idea", "subtitle": "A weekly show about home servers"}
    assert slides[1] == {"layout": "title_content", "title": "Why now", "notes": "Mention the 32 GB home server",
                         "bullets": ["Home servers are cheap", {"text": "A home server is enough", "level": 1}, "People want privacy"]}
    assert slides[2]["layout"] == "title_only" and slides[2]["tables"][0]["rows"] == [["Item", "Price"], ["home server", "₹60,000"]]
    assert slides[3] == {"layout": "section", "title": "Thanks"}


def test_blocks_from_markdown():
    blocks = build.blocks_from_markdown(DOC)
    assert blocks[0] == {"type": "heading", "text": "Trip plan", "level": 1}
    assert blocks[1] == {"type": "paragraph", "runs": [{"text": "Day one in "}, {"text": "Goa", "bold": True}, {"text": ", with "},
                                                      {"text": "sun", "italic": True}, {"text": "."}]}
    assert blocks[2] == {"type": "heading", "text": "Packing", "level": 2}
    assert blocks[3] == {"type": "bullet_list", "items": ["Sunscreen", "Hat"]}
    assert blocks[4] == {"type": "numbered_list", "items": ["Book a cab", "Pack"]}
    assert blocks[5] == {"type": "table", "header": ["Item", "Qty"], "rows": [["Hat", "1"]], "header_bold": True}


def _tools():
    from odie.skills.documents.tools import TOOLS
    return {t.name: t for t in TOOLS}


async def test_make_slides_writes_a_real_pptx(store):
    """Today's Markdown still makes the deck (now designed: odie.decks): titles, bullets with their sub-points,
    notes and the table, in an editable 16:9 PowerPoint. Without a browser (tests), the card is the PowerPoint."""
    out = await _tools()["make_slides"].func(markdown=DECK)
    m = out["media"]
    assert m["kind"] == "file" and m["filename"] == "Our podcast idea.pptx" and (workspace() / m["filename"]).is_file()
    prs = Presentation(str(media.get(m["id"])[0]))

    def texts(slide):
        return " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
    assert [t in texts(s) for t, s in zip(["Our podcast idea", "Why now", "Costs", "Thanks"], prs.slides)] == [True] * 4
    assert "Home servers are cheap" in texts(prs.slides[1]) and "A home server is enough" in texts(prs.slides[1])
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Mention the 32 GB home server"
    assert any(sh.has_table for sh in prs.slides[2].shapes)
    assert prs.slide_width > prs.slide_height  # 16:9
    assert m["deck"]["slides"] == 4 and m["deck"]["theme"] == "editorial"


async def test_make_document_writes_a_real_docx(store):
    out = await _tools()["make_document"].func(markdown=DOC, title="Goa trip")
    m = out["media"]
    assert m["filename"] == "Goa trip.docx" and m["title"] == "Goa trip"
    doc = Document(str(media.get(m["id"])[0]))
    paras = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
    assert ("Heading 1", "Trip plan") in paras and ("Heading 2", "Packing") in paras
    assert ("List Bullet", "Sunscreen") in paras and ("List Number", "Pack") in paras
    assert any(r.bold and r.text == "Goa" for p in doc.paragraphs for r in p.runs)
    assert doc.tables[0].cell(1, 0).text == "Hat"


async def test_bad_markdown_is_a_plain_error(store):
    t = _tools()
    assert "# " in (await t["make_slides"].func(markdown="just some text, no slides"))["error"]
    assert "error" in await t["make_document"].func(markdown="   ")


def test_documents_skill_is_built_in(store):
    from odie.skills import registry
    [e] = [e for e in registry.discover(store) if e.name == "documents"]
    assert e.status != "error" and {"make_slides", "make_document"} <= set(e.tools)
    text = (registry.BUILTIN / "documents" / "SKILL.md").read_text()
    assert "make_slides" in text and "Nous Research" in text


def test_the_hermes_code_keeps_its_licence():
    from odie.skills.documents import hermes
    from pathlib import Path
    folder = Path(hermes.__file__).parent
    lic = (folder / "LICENSE").read_text()
    assert "MIT License" in lic and "Nous Research" in lic and (folder / "NOTICE.md").is_file()


def test_uneven_tables_are_evened():
    md = "# Costs\n| A | B |\n| --- | --- |\n| 1 | 2 | 3 |\n| 4 |\n"
    [slide] = build.slides_from_markdown(md)
    assert slide["tables"][0]["rows"] == [["A", "B"], ["1", "2"], ["4", ""]]
    [table] = build.blocks_from_markdown(md.replace("# Costs\n", ""))
    assert table["rows"] == [["1", "2"], ["4", ""]]
    assert build.build_deck([slide]) and build.build_doc([table])


async def test_tables_in_any_style_and_size_fit(store):
    for sep in ("|-|-|", "| :-: | :-: |", "|:--|--:|"):
        [slide] = build.slides_from_markdown(f"# T\n| A | B |\n{sep}\n| 1 | 2 |\n")
        assert slide["tables"][0]["rows"] == [["A", "B"], ["1", "2"]], sep
    rows = "\n".join(f"| item {i} | {i} |" for i in range(20))
    out = await _tools()["make_slides"].func(markdown=f"# Big\nsub\n\n# Table\n| A | B |\n| --- | --- |\n{rows}\n")
    prs = Presentation(str(media.get(out["media"]["id"])[0]))
    [table] = [sh for sh in prs.slides[1].shapes if sh.has_table]
    assert table.top + table.height <= prs.slide_height


def test_code_blocks_stay_code():
    md = "# Setup\n- Install it\n```bash\n# install it\n- not a bullet\nbrew install uv\n```\n- Done\n"
    slides = build.slides_from_markdown(md)
    assert len(slides) == 1 and slides[0]["bullets"] == ["Install it", "# install it", "- not a bullet", "brew install uv", "Done"]
    blocks = build.blocks_from_markdown("# Setup\n```\n# install it\nbrew install uv\n```\n")
    assert [b["type"] for b in blocks] == ["heading", "paragraph", "paragraph"] and blocks[1]["runs"] == [{"text": "# install it"}]


def test_slide_separators_and_order():
    slides = build.slides_from_markdown("# Title\nsub\n---\n# Two\nIntro line\n- one\n- two\nClosing line\n---\n")
    assert slides[0]["subtitle"] == "sub" and slides[1]["bullets"] == ["Intro line", "one", "two", "Closing line"]


async def test_titles_with_slashes_keep_their_words(store):
    out = await _tools()["make_slides"].func(markdown="# Pros/Cons of a home server\nsub", title="Pros/Cons of a home server")
    assert out["media"]["filename"] == "Pros-Cons of a home server.pptx"


def test_document_quotes_headings_and_images_are_clean():
    blocks = build.blocks_from_markdown("## Head ##\n> quoted words\n\nSee ![a chart](chart.png) here.\n")
    assert blocks[0] == {"type": "heading", "text": "Head", "level": 2}
    assert blocks[1] == {"type": "paragraph", "runs": [{"text": "quoted words", "italic": True}]}
    assert blocks[2] == {"type": "paragraph", "runs": [{"text": "See here."}]}


def test_files_only_when_asked_and_visuals_in_chat():
    from pathlib import Path

    from odie.agent.prompt import Who, rules_text
    rules = rules_text(Who(), gated=True, extra=[])
    assert "with the visual blocks below" in rules and "only when the user asks for a file" in rules
    skill = (Path(__file__).parent.parent / "odie" / "skills" / "documents" / "SKILL.md").read_text()
    assert "Only when the user asks for a deck, slides, a document or a file" in skill
    assert "put a photo on most slides" in skill


@pytest.mark.parametrize("said,deck,doc", [
    ("Make a itinerary from North Carolina for a family of four with food and stay options", False, False),
    ("Make me a beautiful deck for a weekend in Coorg", True, False),
    ("make slides on the quarter", True, False),
    ("put this in a Word document", False, True),
    ("send me a pdf of it", True, True),
    ("same again but in the minimal theme", True, True),
    ("Inline itinerary", False, False),
])
def test_files_only_when_the_user_asks(said, deck, doc):
    from odie.agent import turn
    assert turn.asked_for(said, "deck") is deck and turn.asked_for(said, "document") is doc


async def test_make_slides_bounces_to_chat_when_no_file_was_asked(store):
    from odie.agent import turn
    token = turn.USER_TEXT.set("Make a itinerary for four with food and stays")
    try:
        out = await _tools()["make_slides"].func(markdown=DECK)
        doc = await _tools()["make_document"].func(markdown=DOC)
    finally:
        turn.USER_TEXT.reset(token)
    assert "error" in out and "chat" in out["error"] and "visual blocks" in out["error"]
    assert "error" in doc and "chat" in doc["error"]
    token = turn.USER_TEXT.set("make it a deck")
    try:
        assert "media" in await _tools()["make_slides"].func(markdown=DECK)
    finally:
        turn.USER_TEXT.reset(token)
    assert "media" in await _tools()["make_slides"].func(markdown=DECK)  # no chat turn (background): allowed


def test_stated_facts_are_taken_as_given():
    from odie.agent.prompt import Who, rules_text
    assert "Take what the user states as given" in rules_text(Who(), gated=True, extra=[])
