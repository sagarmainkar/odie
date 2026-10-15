"""The assistant's Markdown as a PowerPoint deck or a Word document. Hermes Agent's builders (hermes/, by Nous
Research, MIT) do the building; this turns Markdown into their specs."""
from __future__ import annotations

import io
import re

from docx import Document
from pptx import Presentation
from pptx.util import Inches

from odie.skills.documents.hermes import docx_create, pptx_create

_IMAGE = re.compile(r"!\[[^\]]*\]\([^)]*\)")
_LINK = re.compile(r"\[([^\]]+)\]\([^)]*\)")
_FENCE = re.compile(r"^\s*(```|~~~)")
_ITEM = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_SEP = re.compile(r"^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)*\|?$")


def _plain(text: str) -> str:
    text = _LINK.sub(r"\1", _IMAGE.sub("", text))
    text = re.sub(r"\*\*|__|`", "", text)
    text = re.sub(r"(?<!\w)[*_](\S[^*_]*?)[*_](?!\w)", r"\1", text)
    return " ".join(text.split())


def _rows(lines: list[str]) -> list[list[str]]:
    """The table's rows, every one as wide as the header (padded or cut): the builders size tables by the first row."""
    rows = [[_plain(c) for c in line.strip().strip("|").split("|")] for line in lines if not _SEP.match(line.strip())]
    width = len(rows[0]) if rows else 0
    return [r[:width] + [""] * (width - len(r)) for r in rows]


def _runs(text: str) -> list[dict]:
    """**bold** and *italic* as runs, the rest plain."""
    text = " ".join(_LINK.sub(r"\1", _IMAGE.sub("", text)).split())
    out = []
    for part in re.split(r"(\*\*[^*]+\*\*|__[^_]+__|(?<!\w)\*[^*\s][^*]*\*(?!\w)|(?<!\w)_[^_\s][^_]*_(?!\w))", text):
        if not part:
            continue
        if part[:2] in ("**", "__"):
            out.append({"text": part[2:-2], "bold": True})
        elif part[0] in "*_" and part[-1] == part[0] and len(part) > 2:
            out.append({"text": part[1:-1], "italic": True})
        else:
            out.append({"text": part.replace("`", "")})
    return out


def slides_from_markdown(md: str) -> list[dict]:
    """"# Title" starts a slide (the first is the title slide, with the line under it as its subtitle); "- " items are
    bullets (two spaces of indent per level); "> " lines are speaker notes; "| a | b |" lines a table; a code block's
    lines are kept as they are; "---" between slides is ignored."""
    slides: list[dict] = []
    cur: dict | None = None
    table: list[str] = []
    fence = False

    def end_table() -> None:
        if table and cur is not None:
            rows = _rows(table)
            if rows:
                height = min(0.45 * len(rows), 5.4)  # stays on the slide: 1.7 in from the top of a 7.5 in slide
                cur.setdefault("tables", []).append({"left": 0.7, "top": 1.7, "width": 11.9, "height": height, "rows": rows})
        table.clear()

    for raw in (md or "").splitlines():
        line = raw.rstrip()
        if _FENCE.match(line):
            fence = not fence
            continue
        if fence:
            if cur is not None and line.strip():
                cur.setdefault("_items", []).append(("code", line.strip()))
            continue
        if re.match(r"^#\s+", line):
            end_table()
            cur = {"title": _plain(line.lstrip("#").rstrip("#"))}
            slides.append(cur)
            continue
        if cur is None:
            continue
        s = line.strip()
        if s.startswith("|"):
            table.append(s)
            continue
        end_table()
        if not s or s in ("---", "***", "___"):
            continue
        if s.startswith(">"):
            cur["notes"] = (cur.get("notes", "") + " " + _plain(s.lstrip("> "))).strip()
        elif m := _ITEM.match(line):
            level = min(4, len(m.group(1).replace("\t", "  ")) // 2)
            text = _plain(m.group(3))
            cur.setdefault("_items", []).append(("bullet", text if level == 0 else {"text": text, "level": level}))
        else:
            cur.setdefault("_items", []).append(("text", _plain(s.lstrip("#"))))
    end_table()
    for i, sl in enumerate(slides):
        items = sl.pop("_items", [])
        if i == 0 and not any(k != "text" for k, _ in items) and not sl.get("tables"):
            sl["layout"] = "title"
            if items:
                sl["subtitle"] = " ".join(v for _, v in items)
        elif items:
            sl["layout"] = "title_content"
            sl["bullets"] = [v for _, v in items]  # in the order they were written
        elif sl.get("tables"):
            sl["layout"] = "title_only"
        else:
            sl["layout"] = "section"
    return slides


def blocks_from_markdown(md: str) -> list[dict]:
    """Headings (#, ##, …), paragraphs with **bold** and *italic*, bullet and numbered lists, tables, "> " quotes in
    italic, code blocks kept as they are, and "---" for a page break."""
    blocks: list[dict] = []
    para: list[str] = []
    items: dict | None = None
    table: list[str] = []
    fence = False

    def end(keep_list: bool = False, keep_table: bool = False) -> None:
        nonlocal items
        if para:
            blocks.append({"type": "paragraph", "runs": _runs(" ".join(para))})
            para.clear()
        if table and not keep_table:
            rows = _rows(table)
            if rows:
                blocks.append({"type": "table", "header": rows[0], "rows": rows[1:], "header_bold": True})
            table.clear()
        if items is not None and not keep_list:
            blocks.append(items)
            items = None

    for raw in (md or "").splitlines():
        s = raw.strip()
        if _FENCE.match(raw):
            end()
            fence = not fence
        elif fence:
            if s:
                blocks.append({"type": "paragraph", "runs": [{"text": raw.rstrip()}]})
        elif not s:
            end()
        elif m := _HEADING.match(s):
            end()
            blocks.append({"type": "heading", "text": _plain(re.sub(r"\s+#+\s*$", "", m.group(2))), "level": min(len(m.group(1)), 4)})
        elif s.startswith("|"):
            end(keep_table=True)
            table.append(s)
        elif s in ("---", "***", "___"):
            end()
            blocks.append({"type": "page_break"})
        elif s.startswith(">"):
            end()
            blocks.append({"type": "paragraph", "runs": [{"text": _plain(s.lstrip("> ")), "italic": True}]})
        elif m := _ITEM.match(raw):
            kind = "numbered_list" if m.group(2)[0].isdigit() else "bullet_list"
            if items is None or items["type"] != kind:
                end()
                items = {"type": kind, "items": []}
            else:
                end(keep_list=True)
            items["items"].append(_plain(m.group(3)))
        else:
            if items is not None:
                end()
            para.append(s)
    end()
    return blocks


def build_deck(slides: list[dict]) -> bytes:
    prs = Presentation()
    ratio = Inches(13.333) / prs.slide_width
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9, as Hermes's builder does
    holders = [*prs.slide_master.placeholders, *(ph for layout in prs.slide_layouts for ph in layout.placeholders)]
    own = [ph for ph in holders if ph._element.xpath("./p:spPr/a:xfrm")]  # inherited ones follow the master
    for holder in own:  # the built-in layouts are 4:3: widen them to fit
        holder.left, holder.width = int(holder.left * ratio), int(holder.width * ratio)
    for spec in slides:
        pptx_create.build_slide(prs, spec)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_doc(blocks: list[dict]) -> bytes:
    doc = Document()
    for block in blocks:
        docx_create.add_block(doc, block)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
