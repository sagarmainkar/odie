"""The native PowerPoint (spec part 4): the same slides, layouts and theme as the HTML deck, built with python-pptx as
real, editable shapes: text boxes in the theme's fonts and colours, photos cropped to fill (rounded where the HTML
rounds them), veils as gradient fills with transparency, rounded cards and tiles, native charts and tables, speaker
notes with the photo credits. Geometry is the HTML's: 1920x1080 px on a 13.333x7.5 in slide (1 px = 6350 EMU), and
type at px/2 points."""
from __future__ import annotations

import io
import math
import re

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from odie.decks.html import _nice, _size  # the same type scale and axis maths as the HTML
from odie.decks.model import Deck, Item, Slide, inline
from odie.decks.photos import Photo
from odie.decks.themes import Theme

PX = 6350
SW, SH, PADX, PADY = 1920, 1080, 120, 100
CW = SW - 2 * PADX  # the content width: 1680
ON_PHOTO, ON_PHOTO_MUTED = "#FFFFFF", "#E3EAF1"


def px(v: float) -> Emu:
    return Emu(int(round(v * PX)))


def _rgba(value: str) -> tuple[str, float]:
    """A CSS colour as (hex without #, alpha 0..1)."""
    m = re.match(r"rgba?\(([\d.]+),\s*([\d.]+),\s*([\d.]+)(?:,\s*([\d.]+))?\)", value)
    if m:
        return "".join(f"{int(float(m[i])):02X}" for i in (1, 2, 3)), float(m[4] or 1)
    return value.lstrip("#").upper()[:6], 1.0


def _flat(value: str, over: str) -> str:
    """A translucent colour flattened over a solid one: the hex that looks the same."""
    top, a = _rgba(value)
    bottom, _ = _rgba(over)
    mix = [round(int(top[i:i + 2], 16) * a + int(bottom[i:i + 2], 16) * (1 - a)) for i in (0, 2, 4)]
    return "".join(f"{v:02X}" for v in mix)


def _rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(_rgba(hex6)[0])


def _clr(value: str) -> etree._Element:
    hex6, a = _rgba(value)
    el = etree.Element(qn("a:srgbClr"), val=hex6)
    if a < 1:
        etree.SubElement(el, qn("a:alpha"), val=str(int(a * 100000)))
    return el


def _set_fill(shape, fill: etree._Element) -> None:
    sp = shape._element.spPr
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill"):
        for el in sp.findall(qn(tag)):
            sp.remove(el)
    geom = sp.find(qn("a:prstGeom"))
    geom.addnext(fill)


def _solid(shape, value: str) -> None:
    f = etree.Element(qn("a:solidFill"))
    f.append(_clr(value))
    _set_fill(shape, f)


def _gradient(shape, stops: list[tuple[float, str]], angle: float) -> None:
    """A linear gradient, angle in degrees (0: left to right, 90: top to bottom), stops (0..1, css colour)."""
    g = etree.Element(qn("a:gradFill"), rotWithShape="1")
    lst = etree.SubElement(g, qn("a:gsLst"))
    for pos, value in stops:
        gs = etree.SubElement(lst, qn("a:gs"), pos=str(int(pos * 100000)))
        gs.append(_clr(value))
    etree.SubElement(g, qn("a:lin"), ang=str(int(angle * 60000)), scaled="0")
    _set_fill(shape, g)


def _no_line(shape) -> None:
    shape.line.fill.background()


def _line(shape, value: str, width_px: float = 1.5) -> None:
    sp = shape._element.spPr
    for el in sp.findall(qn("a:ln")):
        sp.remove(el)
    ln = etree.SubElement(sp, qn("a:ln"), w=str(int(width_px * PX)))
    f = etree.SubElement(ln, qn("a:solidFill"))
    f.append(_clr(value))


def _rect(slide, x, y, w, h, fill: str | None = None, rounded: float = 0.0, line: str | None = None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
    if rounded:
        shape.adjustments[0] = rounded
    shape.shadow.inherit = False
    if fill:
        _solid(shape, fill)
    else:
        shape.fill.background()
    if line:
        _line(shape, line)
    else:
        _no_line(shape)
    return shape


def _lines(text: str, size_px: float, width: float, k: float = 0.52) -> int:
    return max(1, math.ceil(len(text or "") * k * size_px / max(width, 1)))


class _Ctx:
    def __init__(self, t: Theme, photos: dict[str, Photo]):
        self.t, self.photos, self.credits = t, photos, []
        self.bg_mid = _flat("rgba(0,0,0,0)", t.bg)
        self.card = _flat(t.card, t.bg) if t.card.startswith("rgba") else _rgba(t.card)[0]
        self.line = _flat(t.card_line, t.bg) if t.card_line.startswith("rgba") else _rgba(t.card_line)[0]

    def photo(self, ref: str) -> Photo | None:
        p = self.photos.get(ref) if ref else None
        if p is not None and p.credit and p.credit not in self.credits:
            self.credits.append(p.credit)
        return p


def _text(slide, x, y, w, h, content, size_px: float, font: str, color: str, bold: bool = False, italic: bool = False,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing: int = 0, line_spacing: float | None = None, caps: bool = False):
    """A text box: content is a str (with **bold**) or a list of such paragraphs."""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [content]
    for i, text in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        for piece, b in inline(text.upper() if caps else text) or [("", False)]:
            r = p.add_run()
            r.text = piece
            f = r.font
            f.size = Pt(size_px / 2)
            f.name = font
            f.bold = bold or b
            f.italic = italic
            f.color.rgb = _rgb(color)
            if spacing:
                r._r.get_or_add_rPr().set("spc", str(spacing))
    return tb


def _picture(slide, photo: Photo, x, y, w, h, rounded: bool = False):
    pic = slide.shapes.add_picture(str(photo.path), px(x), px(y), px(w), px(h))
    box, img = w / h, photo.width / max(photo.height, 1)
    if img > box:
        c = (1 - box / img) / 2
        pic.crop_left = pic.crop_right = c
    elif img < box:
        c = (1 - img / box) / 2
        pic.crop_top = pic.crop_bottom = c
    if rounded:
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        geom.set("prst", "roundRect")
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(geom, qn("a:avLst"))
        etree.SubElement(av, qn("a:gd"), name="adj", fmla="val 5000")
    return pic


def _background(slide, t: Theme) -> None:
    fill = slide.background.fill
    fill.gradient()
    fill.gradient_angle = 45
    stops = fill.gradient_stops
    stops[0].color.rgb, stops[0].position = _rgb(t.bg), 0
    stops[1].color.rgb, stops[1].position = _rgb(t.bg2), 1.0


def _full_photo(slide, p: Photo, t: Theme, veil: str) -> None:
    _picture(slide, p, 0, 0, SW, SH)
    v = _rect(slide, 0, 0, SW, SH)
    mid = _mid(t.veil_a, t.veil_b)
    if veil == "left":
        _gradient(v, [(0, t.veil_a), (0.46, mid), (1, t.veil_b)], 0)
    elif veil == "bottom":
        _gradient(v, [(0, t.veil_b), (0.5, mid), (1, t.veil_a)], 90)
    else:
        _solid(v, t.veil_a)


def _mid(a: str, b: str, at: float = 0.55) -> str:
    (ha, aa), (hb, ab) = _rgba(a), _rgba(b)
    c = [round(int(ha[i:i + 2], 16) + (int(hb[i:i + 2], 16) - int(ha[i:i + 2], 16)) * at) for i in (0, 2, 4)]
    return f"rgba({c[0]},{c[1]},{c[2]},{aa + (ab - aa) * at:.2f})"


def _head(slide, s: Slide, ctx: _Ctx, y: float = PADY, width: float = CW, x: float = PADX, ink: str | None = None) -> float:
    """Kicker and title; returns where the body starts."""
    t = ctx.t
    if s.kicker:
        _text(slide, x, y, width, 34, s.kicker, 22, t.pptx_body, t.accent, bold=True, spacing=600, caps=True)
        y += 56
    size = _size(s.title, ((28, 76), (44, 64), (80, 54)))
    n = _lines(s.title, size, width, 0.5)
    _text(slide, x, y, width, size * 1.15 * n + 10, s.title, size, t.pptx_head, ink or t.ink, bold=True, line_spacing=1.0)
    y += size * 1.15 * n + 24
    if s.subtitle:
        k = _lines(s.subtitle, 30, min(width, 1400))
        _text(slide, x, y, min(width, 1400), 46 * k, s.subtitle, 30, t.pptx_body, ON_PHOTO_MUTED if ink == ON_PHOTO else t.muted)
        y += 46 * k + 10
    return y + 30


def _bullets(slide, items: list[Item], ctx: _Ctx, x: float, top: float, bottom: float, width: float, on_photo: bool = False,
             two_ok: bool = True, big: bool = False) -> None:
    t = ctx.t
    ink = ON_PHOTO if on_photo else t.ink
    muted = ON_PHOTO_MUTED if on_photo else t.muted
    line = "rgba(255,255,255,0.22)" if on_photo else "#" + ctx.line
    statements = big and len(items) <= 3 and all(it.lead and not it.sub for it in items)
    cols = 2 if two_ok and len(items) > 4 and not statements else 1
    colw = (width - 72 * (cols - 1)) / cols
    size = 30 if len(items) <= 4 else 26
    if statements:
        rows = []
        for it in items:
            h = 46 * 1.2 * _lines(it.lead, 46, colw - 48) + 30 * 1.45 * _lines(it.text, 30, colw - 48) + 60
            rows.append(h)
    else:
        rows = [size * 1.45 * _lines(f"{it.lead} — {it.text}" if it.lead else it.text, size, colw - 44)
                + sum(size * 0.82 * 1.4 for _ in it.sub) + 44 for it in items]
    per_col = math.ceil(len(items) / cols)
    col_h = max(sum(rows[c * per_col:(c + 1) * per_col]) for c in range(cols))
    y0 = top + max(0, (bottom - top - col_h) / 2)
    for i, it in enumerate(items):
        c, r = divmod(i, per_col)
        cx = x + c * (colw + 72)
        cy = y0 + sum(rows[c * per_col:c * per_col + r])
        h = rows[i]
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, px(cx), px(cy + (30 if statements else 22) + size * 0.35), px(14), px(14))
        _solid(dot, t.accent)
        _no_line(dot)
        tx = cx + (48 if statements else 44)
        if statements:
            lh = 46 * 1.2 * _lines(it.lead, 46, colw - 48)
            _text(slide, tx, cy + 22, colw - 48, lh, it.lead, 46, t.pptx_head, ink, bold=True)
            _text(slide, tx, cy + 22 + lh + 6, colw - 48, h - lh - 40, it.text, 30, t.pptx_body, muted)
        else:
            paras = [f"**{it.lead}** — {it.text}" if it.lead and it.text else (f"**{it.lead}**" if it.lead else it.text)]
            tb = _text(slide, tx, cy + 18, colw - 44, h - 30, paras, size, t.pptx_body, muted, line_spacing=1.1)
            for p in tb.text_frame.paragraphs:
                for run in p.runs:
                    if run.font.bold:
                        run.font.color.rgb = _rgb(ink)
            for sub in it.sub:
                q = tb.text_frame.add_paragraph()
                r = q.add_run()
                r.text = f"– {sub}"
                r.font.size, r.font.name, r.font.color.rgb = Pt(size * 0.82 / 2), t.pptx_body, _rgb(muted)
        rule = _rect(slide, cx, cy + h - 1, colw, 1.5)
        _solid(rule, line)


def _cover(slide, s: Slide, ctx: _Ctx) -> None:
    t = ctx.t
    p = ctx.photo(s.image)
    if p:
        _full_photo(slide, p, t, "bottom")
    ink = ON_PHOTO if p else t.ink
    size = _size(s.title, ((12, 168), (20, 140), (32, 112), (52, 92), (80, 78)))
    width = 1500
    y = SH - PADY - 40
    if s.subtitle:
        k = _lines(s.subtitle, 32, 1250)
        y -= 50 * k
        _text(slide, PADX, y, 1250, 50 * k, s.subtitle, 32, t.pptx_body, ink)
    y -= 80
    _rect(slide, PADX, y + 20, 96, 4, fill=t.accent)
    n = _lines(s.title, size, width, 0.55)
    y -= size * 1.0 * n + 10
    _text(slide, PADX, y, width, size * 1.05 * n, s.title, size, t.pptx_head, ink, bold=True, line_spacing=0.95, anchor=MSO_ANCHOR.BOTTOM)
    if s.kicker:
        _text(slide, PADX, y - 50, width, 34, s.kicker, 22, t.pptx_body, t.accent, bold=True, spacing=600, caps=True)


def _worn(slide, ctx: _Ctx, p: Photo | None) -> tuple[str, str, str, str]:
    """A slide that wears its photo behind a full veil: (ink, muted, tile fill, tile line) for it; the theme's without one."""
    t = ctx.t
    if p is None:
        return t.ink, t.muted, "#" + ctx.card, "#" + ctx.line
    _full_photo(slide, p, t, "full")
    return ON_PHOTO, ON_PHOTO_MUTED, "rgba(255,255,255,0.10)", "rgba(255,255,255,0.22)"


def _section(slide, s: Slide, ctx: _Ctx, i: int, p: Photo | None = None) -> None:
    t = ctx.t
    ink, muted, _, _ = _worn(slide, ctx, p)
    faint = _flat(f"rgba({int(t.accent[1:3], 16)},{int(t.accent[3:5], 16)},{int(t.accent[5:7], 16)},0.09)", t.bg2)
    if p is None:
        _text(slide, 1100, 560, 700, 440, f"{i + 1:02d}", 420, t.pptx_head, faint, bold=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.BOTTOM)
    size = _size(s.title, ((24, 120), (44, 96), (80, 76)))
    n = _lines(s.title, size, 1500, 0.55)
    block = (56 if s.kicker else 0) + size * 1.1 * n + 80 + (50 if s.subtitle else 0)
    y = (SH - block) / 2
    if s.kicker:
        _text(slide, PADX, y, CW, 34, s.kicker, 22, t.pptx_body, t.accent, bold=True, spacing=600, caps=True)
        y += 56
    _text(slide, PADX, y, 1500, size * 1.1 * n, s.title, size, t.pptx_head, ink, bold=True, line_spacing=1.0)
    y += size * 1.1 * n + 30
    _rect(slide, PADX, y, 96, 4, fill=t.accent)
    if s.subtitle:
        _text(slide, PADX, y + 40, 1400, 100, s.subtitle, 34, t.pptx_body, muted)


def _points(slide, s: Slide, ctx: _Ctx) -> None:
    top = _head(slide, s, ctx)
    _bullets(slide, s.items, ctx, PADX, top, SH - PADY, CW, big=True)


def _photo(slide, s: Slide, ctx: _Ctx, p: Photo) -> None:
    _full_photo(slide, p, ctx.t, "left")
    width = CW * 0.58
    top = _head(slide, s, ctx, width=width, ink=ON_PHOTO)
    _bullets(slide, s.items, ctx, PADX, top, SH - PADY, width, on_photo=True, two_ok=False)


def _split(slide, s: Slide, ctx: _Ctx, p: Photo) -> None:
    col = (CW - 88) / 2
    top = _head(slide, s, ctx, width=col)
    _bullets(slide, s.items, ctx, PADX, top, SH - PADY, col, two_ok=False)
    fx = PADX + col + 88
    _picture(slide, p, fx, PADY, col, SH - 2 * PADY, rounded=True)
    if s.caption:
        band = _rect(slide, fx, SH - PADY - 110, col, 110)
        _gradient(band, [(0, "rgba(0,0,0,0)"), (1, "rgba(0,0,0,0.72)")], 90)
        _text(slide, fx + 28, SH - PADY - 62, col - 56, 40, s.caption, 20, ctx.t.pptx_body, ON_PHOTO)


def _cards(slide, s: Slide, ctx: _Ctx) -> None:
    t = ctx.t
    top = _head(slide, s, ctx)
    n = len(s.cards)
    gap = 32
    w = (CW - gap * (n - 1)) / n
    any_photo = any(ctx.photos.get(c.image) for c in s.cards)
    if any_photo:
        name_px, line_px, h = {2: 40, 3: 34, 4: 30}.get(n, 30), {2: 26, 3: 23, 4: 21}.get(n, 21), min(SH - PADY - top, 640)
    else:
        name_px, line_px, h = {2: 52, 3: 42, 4: 36}.get(n, 36), {2: 32, 3: 27, 4: 24}.get(n, 24), min(SH - PADY - top, 520)
    y = top + max(0, (SH - PADY - top - h) / 2)
    for k, c in enumerate(s.cards):
        x = PADX + k * (w + gap)
        p = ctx.photo(c.image)
        paper = bool(p)
        _rect(slide, x, y, w, h, fill=t.paper if paper else "#" + ctx.card, rounded=0.06, line=None if paper else "#" + ctx.line)
        ink = t.card_ink if paper else t.ink
        muted = t.card_muted if paper else t.muted
        ty = y
        if p:
            ph = min(340, h * 0.52)
            _picture(slide, p, x, y, w, ph)
            ty = y + ph
        elif any_photo or not p:
            _rect(slide, x + 24, y, w - 48, 8, fill=t.accent)
        pad = 44 if not any_photo else 34
        ty += pad
        nl = _lines(c.title, name_px, w - 2 * pad, 0.55)
        _text(slide, x + pad, ty, w - 2 * pad, name_px * 1.2 * nl, c.title, name_px, t.pptx_head, ink, bold=True)
        ty += name_px * 1.2 * nl + 10
        if c.meta:
            _text(slide, x + pad, ty, w - 2 * pad, 28, c.meta, 18, t.pptx_body, muted, spacing=200, caps=True)
            ty += 38
        for line in c.lines:
            ln = _lines(line, line_px, w - 2 * pad)
            if ty + line_px * 1.5 * ln > y + h - 12:
                break
            tb = _text(slide, x + pad, ty, w - 2 * pad, line_px * 1.5 * ln, line, line_px, t.pptx_body, ink if paper else muted)
            for run in tb.text_frame.paragraphs[0].runs:
                if run.font.bold:
                    run.font.color.rgb = _rgb(t.accent if paper else t.ink)
            ty += line_px * 1.5 * ln + 6


def _steps(slide, s: Slide, ctx: _Ctx, p: Photo | None = None) -> None:
    t = ctx.t
    ink, muted, fill, line = _worn(slide, ctx, p)
    top = _head(slide, s, ctx, ink=ink)
    n = len(s.steps)
    num, tp, dp, h = (110, 36, 28, 460) if n <= 3 else (92, 31, 24, 460) if n <= 5 else (80, 28, 22, 480)
    h = min(h, SH - PADY - top)
    gap = 32 if n <= 4 else 22
    w = (CW - gap * (n - 1)) / n
    y = top + max(0, (SH - PADY - top - h) / 2)
    for k, st in enumerate(s.steps):
        x = PADX + k * (w + gap)
        _rect(slide, x, y, w, h, fill=fill, rounded=0.06, line=line)
        _text(slide, x + 32, y + 36, w - 64, num * 1.05, str(k + 1), num, t.pptx_head, t.accent, bold=True)
        ty = y + 36 + num * 1.05 + 14
        tl = _lines(st.title, tp, w - 64, 0.55)
        _text(slide, x + 32, ty, w - 64, tp * 1.3 * tl, st.title, tp, t.pptx_body, ink, bold=True)
        ty += tp * 1.3 * tl + 12
        if st.detail:
            _text(slide, x + 32, ty, w - 64, max(40, y + h - ty - 20), st.detail, dp, t.pptx_body, muted, line_spacing=1.1)


def _stats(slide, s: Slide, ctx: _Ctx, p: Photo | None = None) -> None:
    t = ctx.t
    ink, muted, fill, line = _worn(slide, ctx, p)
    top = _head(slide, s, ctx, ink=ink)
    n = len(s.stats)
    gap = 32
    w = (CW - gap * (n - 1)) / n
    longest = max(len(st.value) for st in s.stats)
    vp = max(48, min(120, int((w - 80) / (0.62 * longest))))
    h = 340
    extra = 0
    if s.items:
        extra = 44 + max(26 * 1.45 * _lines(f"{i.lead} — {i.text}", 24, (CW - 64) / min(3, len(s.items))) + 30 for i in s.items)
    y = top + max(0, (SH - PADY - top - h - extra) / 2)
    for k, st in enumerate(s.stats):
        x = PADX + k * (w + gap)
        _rect(slide, x, y, w, h, fill=fill, rounded=0.06, line=line)
        _text(slide, x + 40, y + 56, w - 80, vp * 1.1, st.value, vp, t.pptx_head, t.accent, bold=True)
        _text(slide, x + 40, y + 56 + vp * 1.1 + 18, w - 80, h - vp * 1.1 - 100, st.label, 25, t.pptx_body, muted)
    if s.items:
        cols = min(3, len(s.items))
        cw = (CW - 32 * (cols - 1)) / cols
        for k, it in enumerate(s.items[:6]):
            c, r = k % cols, k // cols
            text = f"**{it.lead}** — {it.text}" if it.lead and it.text else (f"**{it.lead}**" if it.lead else it.text)
            ix, iy = PADX + c * (cw + 32), y + h + 44 + r * 70
            if iy + 60 > SH - PADY + 20:
                break
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, px(ix), px(iy + 12), px(12), px(12))
            _solid(dot, t.accent)
            _no_line(dot)
            _text(slide, ix + 28, iy, cw - 28, 60, text, 24, t.pptx_body, muted)


def _numfmt(unit: str, decimals: int) -> str:
    base = "#,##0" + ("." + "0" * decimals if decimals else "")
    if not unit:
        return base
    if unit in ("%", "x", "×"):
        return f'{base}"{unit}"'
    if unit[0].isalpha():
        return f'{base}" {unit}"'
    return f'"{unit[0]}"{base}' + (f'"{unit[1:]}"' if unit[1:] else "")


def _chart(slide, s: Slide, ctx: _Ctx) -> None:
    t = ctx.t
    top = _head(slide, s, ctx)
    c = s.chart
    data = CategoryChartData()
    data.categories = c["x"]
    for se in c["series"]:
        data.add_series(se["name"] or " ", se["data"])
    kind = XL_CHART_TYPE.LINE_MARKERS if c["type"] == "line" else XL_CHART_TYPE.COLUMN_CLUSTERED
    h = SH - PADY - top
    gf = slide.shapes.add_chart(kind, px(PADX), px(top), px(CW), px(h), data)
    ch = gf.chart
    ch.font.size, ch.font.name, ch.font.color.rgb = Pt(11), t.pptx_body, _rgb(t.muted)
    ch.has_legend = len(c["series"]) > 1
    if ch.has_legend:
        ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.TOP, False
    colors = [t.accent, t.muted, t.ink]
    decimals = 0 if all(float(v).is_integer() for se in c["series"] for v in se["data"]) else 1
    for k, series in enumerate(ch.series):
        col = _rgb(colors[k % len(colors)])
        if c["type"] == "line":
            series.format.line.color.rgb = col
            series.format.line.width = Pt(3)
            series.smooth = False
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = col
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = col
    va = ch.value_axis
    peak = max(max(se["data"]) for se in c["series"])
    top_v, step = _nice(peak * 1.08)
    va.maximum_scale, va.minimum_scale, va.major_unit = top_v, 0, step
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = _rgb("#" + ctx.line)
    va.format.line.fill.background()
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = _numfmt(c.get("unit", ""), decimals), False
    ch.category_axis.format.line.color.rgb = _rgb("#" + ctx.line)
    plot = ch.plots[0]
    if c["type"] != "line":
        plot.gap_width = 60
        if len(c["series"]) == 1 and len(c["x"]) <= 12:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.number_format, dl.number_format_is_linked = _numfmt(c.get("unit", ""), decimals), False
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
            dl.font.size, dl.font.bold, dl.font.color.rgb = Pt(12), True, _rgb(t.ink)


def _table(slide, s: Slide, ctx: _Ctx) -> None:
    t = ctx.t
    top = _head(slide, s, ctx)
    rows, cols = len(s.table), max(len(r) for r in s.table)
    body = rows - 1
    size = 28 if body <= 5 else 24 if body <= 8 else 20 if body <= 12 else 15
    rh = min(size * 2.2, (SH - PADY - top) / rows)
    gf = slide.shapes.add_table(rows, cols, px(PADX), px(top), px(CW), px(rh * rows))
    tbl = gf.table
    for row in tbl.rows:
        row.height = px(rh)
    tbl.first_row = tbl.horz_banding = False
    for r in range(rows):
        for k in range(cols):
            cell = tbl.cell(r, k)
            cell.text = ""
            value = s.table[r][k] if k < len(s.table[r]) else ""
            p = cell.text_frame.paragraphs[0]
            for piece, b in inline(value.upper() if r == 0 else value) or [("", False)]:
                run = p.add_run()
                run.text = piece
                run.font.size = Pt((18 if r == 0 else size) / 2)
                run.font.name = t.pptx_body
                run.font.bold = b or r == 0
                run.font.color.rgb = _rgb(t.accent if r == 0 else t.ink)
            cell.margin_left = cell.margin_right = px(26)
            cell.margin_top = cell.margin_bottom = px(4 if body > 8 else 10)
            if r and r % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb("#" + ctx.card)
            else:
                cell.fill.background()


def _gallery(slide, s: Slide, ctx: _Ctx) -> None:
    frames = [(ctx.photo(ref), cap) for ref, _, cap in s.gallery]
    frames = [(p, cap) for p, cap in frames if p]
    if not frames:
        return _points(slide, s, ctx)
    top = _head(slide, s, ctx)
    n = len(frames)
    gap = 28
    w = (CW - gap * (n - 1)) / n
    h = SH - PADY - top
    for k, (p, cap) in enumerate(frames):
        x = PADX + k * (w + gap)
        _picture(slide, p, x, top, w, h, rounded=True)
        if cap:
            band = _rect(slide, x, top + h - 110, w, 110)
            _gradient(band, [(0, "rgba(0,0,0,0)"), (1, "rgba(0,0,0,0.72)")], 90)
            _text(slide, x + 24, top + h - 62, w - 48, 40, cap, 20, ctx.t.pptx_body, ON_PHOTO)


def _quote(slide, s: Slide, ctx: _Ctx) -> None:
    t = ctx.t
    p = ctx.photo(s.image)
    if p:
        _full_photo(slide, p, t, "full")
    ink = ON_PHOTO if p else t.ink
    size = _size(s.quote, ((80, 80), (140, 64), (400, 52)))
    n = _lines(s.quote, size, 1500, 0.5)
    block = 140 + size * 1.2 * n + (90 if s.quote_by else 0)
    y = (SH - block) / 2
    _text(slide, 210, y, 1500, 160, "“", 240, t.pptx_head, t.accent, align=PP_ALIGN.CENTER)
    y += 140
    _text(slide, 210, y, 1500, size * 1.25 * n, s.quote, size, t.pptx_head, ink, italic=True, align=PP_ALIGN.CENTER, line_spacing=1.1)
    if s.quote_by:
        _text(slide, 210, y + size * 1.25 * n + 40, 1500, 44, f"— {s.quote_by}", 30, t.pptx_body, t.accent, align=PP_ALIGN.CENTER,
              spacing=200, caps=True)


def _closing(slide, s: Slide, ctx: _Ctx) -> None:
    t = ctx.t
    p = ctx.photo(s.image)
    if p:
        _full_photo(slide, p, t, "full")
    ink = ON_PHOTO if p else t.ink
    size = _size(s.title, ((16, 150), (30, 116), (80, 88)))
    n = _lines(s.title, size, 1600, 0.55)
    block = size * 1.1 * n + 80 + (60 if s.subtitle else 0)
    y = (SH - block) / 2
    _text(slide, 160, y, 1600, size * 1.1 * n, s.title, size, t.pptx_head, ink, bold=True, align=PP_ALIGN.CENTER)
    y += size * 1.1 * n + 34
    _rect(slide, SW / 2 - 48, y, 96, 4, fill=t.accent)
    if s.subtitle:
        _text(slide, 160, y + 40, 1600, 60, s.subtitle, 32, t.pptx_body, ON_PHOTO_MUTED if p else t.muted, align=PP_ALIGN.CENTER)


def _chrome(slide, i: int, n: int, ctx: _Ctx, on_photo: bool) -> None:
    color = ON_PHOTO_MUTED if on_photo else ctx.t.muted
    _text(slide, SW - 56 - 300, SH - 40 - 26, 300, 26, f"{i + 1:02d} / {n:02d}", 17, ctx.t.pptx_body, color, align=PP_ALIGN.RIGHT, spacing=400)
    if ctx.credits:
        _text(slide, 56, SH - 40 - 26, 1300, 26, "photo: " + " · ".join(ctx.credits), 15, ctx.t.pptx_body, color, spacing=300, caps=True)


def build_pptx(deck: Deck, theme: Theme, photos: dict[str, Photo]) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(SW), px(SH)
    blank = prs.slide_layouts[6]
    n = len(deck.slides)
    for i, s in enumerate(deck.slides):
        slide = prs.slides.add_slide(blank)
        _background(slide, theme)
        ctx = _Ctx(theme, photos)
        layout = s.layout
        on_photo = False
        if layout == "cover":
            _cover(slide, s, ctx)
            on_photo = bool(photos.get(s.image))
        elif layout in ("photo", "split"):
            p = ctx.photo(s.image)
            if p is None:
                _points(slide, s, ctx)
            elif layout == "photo":
                _photo(slide, s, ctx, p)
                on_photo = True
            else:
                _split(slide, s, ctx, p)
        elif layout in ("section", "steps", "stats"):  # these wear their photo behind a full veil
            p = ctx.photo(s.image)
            {"section": lambda: _section(slide, s, ctx, i, p), "steps": lambda: _steps(slide, s, ctx, p),
             "stats": lambda: _stats(slide, s, ctx, p)}[layout]()
            on_photo = p is not None
        elif layout == "cards":
            _cards(slide, s, ctx)
        elif layout == "table":
            _table(slide, s, ctx)
        elif layout == "chart":
            _chart(slide, s, ctx)
        elif layout == "gallery":
            _gallery(slide, s, ctx)
        elif layout == "quote":
            _quote(slide, s, ctx)
            on_photo = bool(photos.get(s.image))
        elif layout == "closing":
            _closing(slide, s, ctx)
            on_photo = bool(photos.get(s.image))
        else:
            _points(slide, s, ctx)
        _chrome(slide, i, n, ctx, on_photo)
        notes = list(s.notes) + [f"Photo: {c}" for c in ctx.credits]
        if notes:
            slide.notes_slide.notes_text_frame.text = "\n".join(notes)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
